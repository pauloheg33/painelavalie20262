"""Importa a consolidação do AVALIE.CE para SQLite, sem modificar o Excel."""
import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import sqlite3
import tempfile

import openpyxl

from backend.services.data_service import AVALIACAO, DEFAULT_DB
from backend.import_records import read_individual

ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE = ROOT / "backend" / "data" / "DADOS_ACERTO_POR_HABILIDADE.xlsx"
DEFAULT_INDIVIDUAL = ROOT / "backend" / "data" / "Database.xlsx"
BOOKLETS = {2: {"LP": ("P0201", 1, 22), "MT": ("M0201", 1, 22)},
            4: {"LP": ("C0401", 1, 22), "MT": ("C0401", 23, 44)},
            5: {"LP": ("C0501", 1, 22), "MT": ("C0501", 23, 44)},
            8: {"LP": ("C0801", 1, 26), "MT": ("C0801", 27, 52)},
            9: {"LP": ("C0901", 1, 26), "MT": ("C0901", 27, 52)}}

SCHEMA = """
CREATE TABLE schools (id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE);
CREATE TABLE items (
 id TEXT PRIMARY KEY, year INTEGER NOT NULL CHECK(year IN (2,4,5,8,9)),
 component TEXT NOT NULL CHECK(component IN ('LP','MT')), booklet TEXT NOT NULL,
 number INTEGER NOT NULL, description TEXT NOT NULL,
 UNIQUE(year, component, booklet, number)
);
CREATE TABLE item_rates (
 school_id INTEGER NOT NULL REFERENCES schools(id), item_id TEXT NOT NULL REFERENCES items(id),
 rate REAL NOT NULL CHECK(rate BETWEEN 0 AND 1), source_sheet TEXT NOT NULL,
 source_row INTEGER NOT NULL, PRIMARY KEY(school_id,item_id)
);
CREATE TABLE evaluation_records (
 id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), student_name TEXT NOT NULL,
 year INTEGER NOT NULL, booklet TEXT NOT NULL, discipline TEXT NOT NULL,
 correct_count INTEGER NOT NULL, wrong_count INTEGER NOT NULL, reported_rate REAL NOT NULL,
 wrong_items TEXT NOT NULL, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL,
 UNIQUE(source_sheet,source_row)
);
CREATE TABLE item_answers (
 record_id INTEGER NOT NULL REFERENCES evaluation_records(id), item_id TEXT NOT NULL REFERENCES items(id),
 alternative TEXT, is_correct INTEGER NOT NULL CHECK(is_correct IN (0,1)), PRIMARY KEY(record_id,item_id)
);
CREATE INDEX answer_item_index ON item_answers(item_id);
CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
"""


def read_source(source):
    records, definitions, coverage = [], {}, defaultdict(set)
    with source.open("rb") as stream:
        workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            for year, components in BOOKLETS.items():
                sheet_name = f"ITENS_{year}_ANO"
                if sheet_name not in workbook.sheetnames:
                    raise ValueError(f"Aba obrigatória ausente: {sheet_name}")
                sheet = workbook[sheet_name]
                for row_number, cells in enumerate(sheet.iter_rows(min_row=4), 4):
                    values = [c.value for c in cells[:8]]
                    if all(v is None for v in values):
                        continue
                    evaluation, network, grade, component, school, item, description, rate = values
                    location = f"{sheet_name}:{row_number}"
                    if evaluation != AVALIACAO or str(network).strip() != "MUNICIPAL":
                        raise ValueError(f"Avaliação ou rede inesperada em {location}")
                    # O rótulo contém também '9 ANOS': ler o ano no final do texto.
                    match = re.search(r"(\d+)\s*º?\s*ANO\s*$", str(grade), re.I)
                    if not match or int(match[1]) != year or component not in components:
                        raise ValueError(f"Ano/componente inválido em {location}")
                    match = re.fullmatch(r"Item\s+(\d+)", str(item), re.I)
                    booklet, first, last = components[component]
                    if not match or not first <= int(match[1]) <= last:
                        raise ValueError(f"Item fora do caderno {booklet} em {location}")
                    if not school or not description:
                        raise ValueError(f"Escola ou descrição vazia em {location}")
                    if isinstance(rate, bool) or not isinstance(rate, (int, float)) or not math.isfinite(rate) or not 0 <= rate <= 1:
                        raise ValueError(f"Taxa ausente/inválida em {location}; esperado valor de 0 a 1")
                    number = int(match[1])
                    key = f"{year}:{component}:{booklet}:{number:02d}"
                    definition = (key, year, component, booklet, number, str(description).strip())
                    if key in definitions and definitions[key] != definition:
                        raise ValueError(f"Descrições divergentes para {key} em {location}")
                    definitions[key] = definition
                    school = " ".join(str(school).split())
                    group = (school, year, component)
                    if number in coverage[group]:
                        raise ValueError(f"Item duplicado em {location}: {school}, {key}")
                    coverage[group].add(number)
                    records.append((school, key, float(rate), sheet_name, row_number))
        finally:
            workbook.close()
    for (school, year, component), numbers in coverage.items():
        _, first, last = BOOKLETS[year][component]
        if numbers != set(range(first, last + 1)):
            raise ValueError(f"Cobertura incompleta: {school}, {year}, {component}")
    if not records:
        raise ValueError("A consolidação não contém taxas de itens")
    return records, definitions


def build(source=DEFAULT_SOURCE, database=DEFAULT_DB, individual=DEFAULT_INDIVIDUAL):
    source, database = Path(source).resolve(), Path(database).resolve()
    if source == database:
        raise ValueError("O banco de destino não pode substituir a planilha de origem")
    records, items = read_source(source)
    schools = {name: idx for idx, name in enumerate(sorted({r[0] for r in records}), 1)}
    individual = Path(individual).resolve() if individual else None
    evaluations, answers, issues = [], [], []
    if individual is not None:
        if individual == database:
            raise ValueError('O banco não pode substituir a base individual')
        evaluations, answers, counts, successes, issues = read_individual(individual, schools, BOOKLETS)
        source_keys = {(r[0], r[1]) for r in records}
        if source_keys != set(counts):
            raise ValueError('Cobertura da base individual diverge da consolidação')
        differences = [(school, key) for school, key, rate, _, _ in records
                       if abs(rate - successes[school,key] / counts[school,key]) > 1e-9]
        if differences:
            raise ValueError(f'{len(differences)} taxas divergem das listas de questões erradas: {differences[:3]}')
    metadata = {"avaliacao": AVALIACAO, "source": source.name,
                "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "imported_at": datetime.now(timezone.utc).isoformat(),
                "taxas_por_escola_item": len(records), "itens_distintos": len(items),
                "escolas": len(schools), "base_individual_importada": bool(evaluations),
                "registros_avaliacao": len(evaluations) if evaluations else None,
                "estudantes_distintos": None, "schema_version": 2,
                "taxas_reconciliadas": len(records) if evaluations else 0}
    metadata['divergencias_totais_listas'] = issues
    if evaluations:
        metadata['individual_source'] = individual.name
        metadata['individual_sha256'] = hashlib.sha256(individual.read_bytes()).hexdigest()
    database.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=database.parent, suffix=".sqlite3", delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        with sqlite3.connect(temp_path) as conn:
            conn.execute("PRAGMA foreign_keys=ON")
            conn.executescript(SCHEMA)
            conn.executemany("INSERT INTO schools VALUES (?,?)", [(idx, name) for name, idx in schools.items()])
            conn.executemany("INSERT INTO items VALUES (?,?,?,?,?,?)", items.values())
            conn.executemany("INSERT INTO item_rates VALUES (?,?,?,?,?)",
                             [(schools[s], key, rate, sheet, row) for s, key, rate, sheet, row in records])
            conn.executemany('INSERT INTO evaluation_records VALUES (?,?,?,?,?,?,?,?,?,?,?,?)', evaluations)
            conn.executemany('INSERT INTO item_answers VALUES (?,?,?,?)', answers)
            conn.executemany("INSERT INTO metadata VALUES (?,?)", [(k, json.dumps(v)) for k, v in metadata.items()])
        conn.close()
        temp_path.replace(database)
    finally:
        temp_path.unlink(missing_ok=True)
    print(f"Importação OK: {len(records)} taxas, {len(items)} itens, {len(schools)} escolas. Banco: {database}")
    return metadata


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--database", type=Path, default=DEFAULT_DB)
    parser.add_argument("--individual", type=Path, default=DEFAULT_INDIVIDUAL)
    parser.add_argument("--consolidated-only", action="store_true")
    args = parser.parse_args()
    build(args.source, args.database, None if args.consolidated_only else args.individual)
