"""Leitura local dos registros individuais e apuração pelas questões erradas."""
from collections import Counter
import math
import re
import unicodedata

import openpyxl


def school_key(name):
    value = ''.join(c for c in unicodedata.normalize('NFD', str(name).upper())
                    if not unicodedata.combining(c))
    return re.sub(r'^EE(?:I)?EF\s+|^EEF\s+', '', ' '.join(value.split()))


def parse_wrong_items(value, total, errors):
    if str(value).strip().casefold() == 'nenhuma' and errors == 0:
        return set()
    if value is None or not str(value).strip():
        if errors != 0:
            raise ValueError('Lista de questões erradas vazia com erros informados')
        return set()
    text = str(value).strip()
    if not re.fullmatch(r'\d+(?:\s*,\s*\d+)*', text):
        raise ValueError('Lista de questões erradas com formato inválido')
    numbers = [int(n.strip()) for n in text.split(',')]
    if len(numbers) != len(set(numbers)):
        raise ValueError('Lista de questões erradas contém duplicatas')
    if any(n < 1 or n > total for n in numbers):
        raise ValueError('Questão errada fora dos limites do caderno')
    return set(numbers)


def read_individual(source, schools, booklets):
    lookup = {school_key(name): name for name in schools}
    # Variante abreviada presente na base individual; identificação oficial na consolidação.
    lookup['MARIA AMELIA RODRIGUES'] = 'EEIEF MARIA AMÉLIA RODRIGUES DE SOUSA'
    records, answers, issues = [], [], []
    counts = Counter()
    successes = Counter()
    with source.open('rb') as stream:
        workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            for sheet_name in ['2º ano', '4º e 5º', '8º e 9º']:
                if sheet_name not in workbook.sheetnames:
                    raise ValueError(f'Aba individual ausente: {sheet_name}')
                rows = workbook[sheet_name].iter_rows(values_only=True)
                header = next(rows)
                columns = {label: idx for idx, label in enumerate(header) if label is not None}
                required = ['Nome do aluno', 'Escola', 'Caderno', 'Disciplina', 'Ano',
                            'Acertos', 'Erros', '%', 'Questões erradas']
                if any(c not in columns for c in required):
                    raise ValueError(f'Colunas individuais ausentes em {sheet_name}')
                for row_number, row in enumerate(rows, 2):
                    if all(v is None for v in row):
                        continue
                    get = lambda key: row[columns[key]]
                    location = f'{sheet_name}:{row_number}'
                    try:
                        if not get('Nome do aluno'):
                            raise ValueError('Nome do aluno ausente')
                        school = lookup.get(school_key(get('Escola')))
                        if school not in schools:
                            raise ValueError('Escola sem correspondência na consolidação')
                        match = re.fullmatch(r'(\d+)º?\s*ano', str(get('Ano')).strip(), re.I)
                        if not match or int(match[1]) not in booklets:
                            raise ValueError('Ano inválido')
                        year = int(match[1])
                        booklet = get('Caderno')
                        components = {c: spec for c, spec in booklets[year].items() if spec[0] == booklet}
                        if not components:
                            raise ValueError('Caderno incompatível com o ano')
                        expected_discipline = ('Língua Portuguesa' if 'LP' in components else 'Matemática') if year == 2 else 'Língua Portuguesa e Matemática'
                        if get('Disciplina') != expected_discipline:
                            raise ValueError('Disciplina incompatível com o caderno')
                        total = max(spec[2] for spec in components.values())
                        correct, errors = get('Acertos'), get('Erros')
                        if any(isinstance(v, bool) or not isinstance(v, (int,float)) or not math.isfinite(v) or v != int(v) or v < 0 for v in (correct, errors)) or correct + errors != total:
                            raise ValueError('Acertos e erros incompatíveis com o total de questões')
                        wrong = parse_wrong_items(get('Questões erradas'), total, errors)
                        if len(wrong) != errors:
                            issues.append({'aba': sheet_name, 'linha': row_number, 'escola': school,
                                           'caderno': booklet, 'erros_informados': int(errors),
                                           'questoes_na_lista': len(wrong)})
                        reported = get('%')
                        if isinstance(reported, str):
                            reported = float(reported.strip().replace('%','').replace(',','.')) / 100
                        if not isinstance(reported, (int,float)) or not math.isfinite(reported) or abs(reported - correct / total) > 0.00051:
                            raise ValueError('Percentual geral diverge dos acertos')
                        record_id = len(records) + 1
                        records.append((record_id, schools[school], str(get('Nome do aluno')).strip(), year,
                                        booklet, get('Disciplina'), int(correct), int(errors), reported,
                                        str(get('Questões erradas') or ''), sheet_name, row_number))
                        for component, (_, first, last) in components.items():
                            for number in range(first, last + 1):
                                if number not in columns:
                                    raise ValueError(f'Coluna do item {number} ausente')
                                key = f'{year}:{component}:{booklet}:{number:02d}'
                                success = int(number not in wrong)
                                alternative = get(number)
                                answers.append((record_id, key, None if alternative is None else str(alternative), success))
                                counts[school, key] += 1
                                successes[school, key] += success
                    except (ValueError, TypeError) as exc:
                        raise ValueError(f'{location}: {exc}') from exc
        finally:
            workbook.close()
    return records, answers, counts, successes, issues
